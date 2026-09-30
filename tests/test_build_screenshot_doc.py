"""Tests for tools/build_screenshot_doc.py."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import build_screenshot_doc  # noqa: E402


def test_missing_folder_gives_placeholders(tmp_path: Path) -> None:
    out, real, missing = build_screenshot_doc.build(tmp_path / "absent", tmp_path / "doc.docx")
    assert (real, missing) == (0, 14)
    text = "\n".join(p.text for t in Document(out).tables for c in t._cells for p in c.paragraphs)
    assert text.count("PLACEHOLDER") == 14


def test_images_sorted_naturally_and_embedded(tmp_path: Path) -> None:
    source = tmp_path / "shots"
    source.mkdir()
    for name in ["10_last.png", "2_second.png", "1_first_loop.png", "notes.txt"]:
        path = source / name
        if path.suffix == ".png":
            Image.new("RGB", (400, 300), "white").save(path)
        else:
            path.write_text("ignored", encoding="utf-8")
    images = build_screenshot_doc.find_images(source)
    assert [p.name for p in images] == ["1_first_loop.png", "2_second.png", "10_last.png"]
    assert build_screenshot_doc.caption_for(images[0], 1) == "Exercise 01 — first loop"

    out, real, missing = build_screenshot_doc.build(source, tmp_path / "doc.docx")
    assert (real, missing) == (3, 11)
    assert len(Document(out).inline_shapes) == 3


def test_numbered_files_keep_their_slot(tmp_path: Path) -> None:
    # Regression: with 02 missing, 03_loops.png used to be captioned "Exercise 02".
    source = tmp_path / "shots"
    source.mkdir()
    for name in ["01_sum.png", "03_loops.png", "extra.png"]:
        Image.new("RGB", (200, 100), "white").save(source / name)
    slots = build_screenshot_doc.assign_slots(build_screenshot_doc.find_images(source))
    assert {n: p.name for n, p in slots.items()} == {
        1: "01_sum.png",
        3: "03_loops.png",
        2: "extra.png",
    }
