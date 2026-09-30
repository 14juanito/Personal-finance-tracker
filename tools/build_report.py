"""Build the project report (deliverables/Project_Report.docx).

Steps:
    1. Draw the architecture diagram (docs/architecture.png) with matplotlib.
    2. Run the test suite with coverage to quote real, current numbers.
    3. Write the report with python-docx, embedding the real screenshots from
       docs/screenshots/ (produced by tools/capture_screenshots.py).

Usage:
    python tools/build_report.py [--skip-tests]

Main elements:
    draw_architecture, test_stats, docx helpers (heading, para, table, figure), build.

Course concepts illustrated:
    Functions, lists of rows, subprocess + JSON parsing for real test figures.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from matplotlib.figure import Figure
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots"
ARCHITECTURE_PNG = ROOT / "docs" / "architecture.png"
OUTPUT = ROOT / "deliverables" / "Project_Report.docx"
ACCENT = RGBColor(0x15, 0x65, 0xC0)
GREY = RGBColor(0x61, 0x61, 0x61)
HEADER_FILL = "1565C0"
ROW_FILL = "EEF3FB"
TABLE_WIDTH_IN = 6.5


# --------------------------------------------------------------------------- diagram
def draw_architecture(path: Path = ARCHITECTURE_PNG) -> Path:
    """Draw the layered architecture diagram and save it as PNG.

    Args:
        path: Destination image.

    Returns:
        The path written.
    """
    fig = Figure(figsize=(10, 6.2), dpi=160)
    ax = fig.add_axes((0, 0, 1, 1))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 62)
    ax.axis("off")

    layers = [
        (
            50,
            "User interfaces",
            "#E3F2FD",
            [
                "main.py\n(cli | gui | demo)",
                "cli.py\nconsole menu",
                "gui_tkinter.py\nTkinter window",
                "dashboard.py\nStreamlit",
            ],
        ),
        (
            35,
            "Application logic",
            "#E8F5E9",
            [
                "tracker.py\nFinanceTracker",
                "alerts.py\nbudget alerts",
                "analytics.py\npandas",
                "visualize.py\nmatplotlib",
            ],
        ),
        (
            20,
            "Domain & persistence",
            "#FFF3E0",
            [
                "models.py\ndataclasses",
                "exceptions.py\nerror classes",
                "storage.py\nJSON + CSV",
                "config.py\npaths",
            ],
        ),
        (
            5,
            "Files",
            "#F3E5F5",
            ["data/sample_data.json", "data/*.csv", "output/charts/*.png", "data/my_finances.json"],
        ),
    ]
    for y, title, color, boxes in layers:
        ax.add_patch(
            FancyBboxPatch((1, y - 1), 98, 12, boxstyle="round,pad=0.3", fc=color, ec="#BDBDBD")
        )
        ax.text(2.5, y + 9.7, title, fontsize=10, weight="bold", color="#424242")
        for index, label in enumerate(boxes):
            x = 5 + index * 23.5
            ax.add_patch(
                FancyBboxPatch(
                    (x, y + 0.5), 19, 7, boxstyle="round,pad=0.4", fc="white", ec="#757575"
                )
            )
            ax.text(x + 9.5, y + 4, label, ha="center", va="center", fontsize=8.5)

    # One arrow between consecutive layers: each layer only uses the layer below it,
    # which reads more clearly than a web of module-to-module arrows.
    for upper, lower in zip(layers, layers[1:], strict=False):
        y_top, y_bottom = upper[0] - 1.3, lower[0] + 11.3
        for x in (27, 73):
            ax.add_patch(
                FancyArrowPatch(
                    (x, y_top),
                    (x, y_bottom),
                    arrowstyle="-|>",
                    mutation_scale=16,
                    color="#546E7A",
                    lw=1.6,
                )
            )
        ax.text(
            50,
            (y_top + y_bottom) / 2,
            "uses",
            ha="center",
            va="center",
            fontsize=8,
            color="#546E7A",
            style="italic",
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    return path


# --------------------------------------------------------------------------- test stats
def test_stats(skip: bool) -> dict[str, str]:
    """Run pytest with coverage and return the numbers quoted in the report.

    Args:
        skip: Do not run the tests (placeholder markers are returned instead).

    Returns:
        ``{"tests": count, "coverage": percent}`` as strings.

    Raises:
        SystemExit: If the test suite fails.
    """
    if skip:
        # Visible markers instead of invented numbers: a draft built without tests
        # must never look like it quotes real results.
        return {"tests": "[tests not run]", "coverage": "[coverage not measured]"}
    with tempfile.TemporaryDirectory() as tmp:
        report = Path(tmp) / "coverage.json"
        result = subprocess.run(
            [sys.executable, "-m", "pytest", "-q", "--cov", f"--cov-report=json:{report}"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            raise SystemExit("Tests failed — fix them before building the report.")
        passed = re.search(r"(\d+) passed", result.stdout)
        coverage = json.loads(report.read_text(encoding="utf-8"))["totals"]["percent_covered"]
    return {"tests": passed.group(1) if passed else "?", "coverage": f"{coverage:.0f}"}


# --------------------------------------------------------------------------- docx helpers
def shade(cell: _Cell, fill: str) -> None:
    """Give a table cell a background colour (clear shading, never solid).

    Args:
        cell: Table cell.
        fill: Hex colour without ``#``.
    """
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def heading(doc: Document, text: str, level: int = 1) -> None:
    """Add a numbered section heading in the accent colour.

    Args:
        doc: Document being built.
        text: Heading text.
        level: Heading level (1 = section).
    """
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = ACCENT


def add_rich_text(
    paragraph: Paragraph, text: str, italic: bool = False, size: int | None = None
) -> None:
    """Add text where **bold** and `code` segments are formatted (markers removed).

    Args:
        paragraph: Paragraph receiving the runs.
        text: Text with optional ``**bold**`` and backtick-quoted code segments.
        italic: Make every run italic.
        size: Font size in points (None = document default).
    """
    for chunk in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not chunk:
            continue
        if chunk.startswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(chunk)
        run.italic = italic
        if size:
            run.font.size = Pt(size)


def para(doc: Document, text: str, italic: bool = False, size: int | None = None) -> None:
    """Add a body paragraph with inline formatting.

    Args:
        doc: Document being built.
        text: Paragraph text (see ``add_rich_text``).
        italic: Italic paragraph.
        size: Font size in points.
    """
    add_rich_text(doc.add_paragraph(), text, italic, size)


def bullets(doc: Document, items: list[str]) -> None:
    """Add a bulleted list using Word's built-in list style.

    Args:
        doc: Document being built.
        items: One string per bullet.
    """
    for item in items:
        add_rich_text(doc.add_paragraph(style="List Bullet"), item)


def figure(doc: Document, path: Path, caption: str, width_in: float = 6.0) -> None:
    """Add a centred image followed by an italic caption.

    Args:
        doc: Document being built.
        path: Image file.
        caption: Caption text.
        width_in: Image width in inches.
    """
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic, run.font.size, run.font.color.rgb = True, Pt(9), GREY


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    """Add a table with explicit column widths summing to the text width.

    Args:
        doc: Document being built.
        headers: Column titles.
        rows: Cell texts (backtick-quoted parts are set in a code font).
        widths: Column widths in inches.
    """
    assert abs(sum(widths) - TABLE_WIDTH_IN) < 0.01
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    for index, text in enumerate(headers):
        cell = t.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.size = True, Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, HEADER_FILL)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = ""
            parts = re.split(r"`(.+?)`", text)
            for part_index, part in enumerate(parts):
                run = cells[index].paragraphs[0].add_run(part)
                run.font.size = Pt(9)
                if part_index % 2 == 1:
                    run.font.name = "Consolas"
            if row_index % 2 == 1:
                shade(cells[index], ROW_FILL)
    for row in t.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
    doc.add_paragraph()


# --------------------------------------------------------------------------- content
CONCEPTS: list[list[str]] = [
    [
        "User input / output",
        "`cli.py` → `ConsoleApp.ask`, `prompt_*`, `show_summary`",
        "All keyboard input goes through one method; results are printed as aligned f-string tables.",
    ],
    [
        "Decision structures",
        "`alerts.budget_level`, `main.main`, `ConsoleApp.run`",
        "`if/elif/else` classifies budget usage (OK / WARNING / EXCEEDED) and dispatches the modes.",
    ],
    [
        "Loops",
        "`ConsoleApp.run`, `prompt_amount`, `FinanceTracker.filter`",
        "`while` menu loop, re-prompt loops for invalid input, `for` loops with `continue` filters.",
    ],
    [
        "Functions",
        "`models.parse_amount`, `analytics.*`, `visualize.*`",
        "Small, documented, type-hinted functions reused by the CLI, GUI and dashboard.",
    ],
    [
        "File handling",
        "`storage.save_json/load_json/save_csv/load_csv`, `config.py`",
        "`pathlib`, `with open(...)`, `json` and `csv` modules, atomic writes via a temp file.",
    ],
    [
        "Exceptions",
        "`exceptions.py`, `storage.load_json`, `models.parse_date`",
        "Custom hierarchy; `try/except/else/finally`; corrupted JSON is backed up and recovered.",
    ],
    [
        "Lists",
        "`FinanceTracker.transactions`, list comprehensions",
        "Ordered collection of transactions; comprehensions build filtered views.",
    ],
    [
        "Dictionaries",
        "`FinanceTracker.budgets/goals`, `ConsoleApp.menu`, `to_dict()`",
        "O(1) lookup by name, menu dispatch table, JSON serialization.",
    ],
    [
        "Sets",
        "`FinanceTracker.categories`, `models.VALID_KINDS`, `storage.load_csv`",
        "Unique category names, O(1) membership tests, set difference for missing CSV columns.",
    ],
    [
        "OOP (classes)",
        "`Transaction`, `SavingsGoal`, `Budget`, `FinanceTracker`, `FinanceApp`",
        "Dataclasses validated in `__post_init__`, properties, class methods, inheritance.",
    ],
    [
        "pandas",
        "`analytics.monthly_summary`, `category_breakdown`, `spending_trend`",
        "`pivot_table`, `groupby().agg()`, `rolling().mean()`, `pct_change()`, boolean masks.",
    ],
]

CHALLENGES: list[list[str]] = [
    [
        "Our own `StorageError` (a subclass of `OSError`) was caught by a generic `except OSError` "
        "and re-wrapped, hiding the real message.",
        "Added `except StorageError: raise` before the generic clause — handlers are "
        "checked in order.",
    ],
    [
        'Goal labels such as "$1,850 / $3,000" lost their dollar signs: matplotlib treats text '
        "between two `$` as LaTeX math. The same happened in Streamlit's `st.markdown`.",
        "Escaped the `$` signs (`\\$`) in both places and added regression tests.",
    ],
    [
        "CLI tests tried to read the real keyboard because `input_func=input` was a default "
        "argument, evaluated once when the function is defined.",
        "Defaults changed to `None`, resolved at call time; tests inject a scripted input function.",
    ],
    [
        "On a HiDPI screen the Tkinter table rows overlapped and chart text was tiny: Tk scales "
        "fonts but not pixel sizes, and the canvas resets a figure's dpi.",
        "Computed a scale factor from the real font height; added a `dpi` parameter to the charts.",
    ],
    [
        'First sample data had a negative savings rate and all goals "overdue" '
        "(data ended in 2025).",
        "Re-tuned the generator (fixed seed 42) to Apr–Sep 2026 and added a test on the balance.",
    ],
    [
        "Streamlit 1.64 deprecated `use_container_width`, flooding the log with warnings.",
        'Switched to the current `width="stretch"` API and verified a clean server log.',
    ],
]


def build(skip_tests: bool = False, output: Path = OUTPUT) -> Path:
    """Write the report.

    Args:
        skip_tests: Do not run pytest (the test figures then show "not run" markers).
        output: Destination ``.docx`` file.

    Returns:
        The path written.
    """
    stats = test_stats(skip_tests)
    draw_architecture()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    # ---------------------------------------------------------------- title page
    for _ in range(7):
        doc.add_paragraph()
    for text, size, bold, color in [
        ("Personal Finance Tracker", 30, True, ACCENT),
        (
            "A Python application for tracking, analysing and visualising personal finances",
            14,
            False,
            GREY,
        ),
        ("", 12, False, None),
        ("DATA 333 – Data Management & Analysis", 13, True, None),
        ("Bellevue College — Prior Learning Assessment (PLA)", 12, False, None),
        ("Project Report", 12, False, None),
        ("", 12, False, None),
        ("[Student Name]", 12, False, None),
        ("[Date]", 12, False, None),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size, run.bold = Pt(size), bold
        if color:
            run.font.color.rgb = color
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- 1. overview
    heading(doc, "1. Project Overview and Objectives")
    para(
        doc,
        "The Personal Finance Tracker is a Python application that records income and "
        "expenses, organises them into categories, and turns them into summaries, trends, "
        "budget alerts and charts. It offers three ways to work with the same data: an "
        "interactive **console menu**, a **Tkinter desktop window** and an interactive "
        "**Streamlit web dashboard**, plus a non-interactive **demo mode** that runs every "
        "feature on a realistic sample dataset (311 fictional transactions over six months).",
    )
    para(doc, "The objectives, taken from the DATA 333 project brief, were to:")
    bullets(
        doc,
        [
            "track income and expenses and assign each transaction to a category;",
            "display spending summaries and analyse spending trends over time;",
            "set savings goals and follow their progress;",
            "save and load data from files (both **JSON** and **CSV**);",
            "visualise the results, and demonstrate every core course concept: input/output, "
            "decisions and loops, functions, files and exceptions, lists/dictionaries/sets, "
            "object-oriented programming and pandas;",
            "implement all four enhancements: a **Tkinter GUI**, **JSON + CSV storage**, "
            "**budget alerts** and a **Streamlit dashboard**.",
        ],
    )

    # ---------------------------------------------------------------- 2. features
    heading(doc, "2. Key Features")
    bullets(
        doc,
        [
            "**Transactions** — add, edit, delete, search (text) and filter (kind, category, "
            "date range, amount). Every value is validated; bad input is explained, never fatal.",
            "**Summaries & analysis** — monthly income/expenses/net and savings rate, category "
            "shares, 3-month rolling average, month-over-month change and top expenses (pandas).",
            "**Budgets & alerts** — a monthly limit per category; WARNING at 80 %, EXCEEDED at "
            "100 %. The CLI and GUI warn immediately when a new expense crosses a threshold.",
            "**Savings goals** — target, deadline, contributions, milestones at 25/50/75/100 %.",
            "**Files** — full state in JSON, transactions in CSV (import/export). Writes are "
            "atomic; a corrupted JSON file is backed up and the app starts cleanly.",
            "**Charts** — category donut, monthly income vs. expenses, spending trend and goal "
            "progress, saved as PNG and embedded in the GUI; interactive Plotly charts online.",
        ],
    )
    figure(
        doc,
        SHOTS / "cli_session.png",
        "Figure 1 — CLI mode: an invalid amount is re-prompted, then a budget alert is raised.",
        3.4,
    )
    figure(
        doc,
        SHOTS / "gui_summary.png",
        "Figure 2 — Tkinter GUI, Summary tab: KPIs, monthly table and colour-coded alerts.",
        5.6,
    )
    figure(
        doc,
        SHOTS / "dashboard_overview.png",
        "Figure 3 — Streamlit dashboard (captured with Playwright): filters, KPIs, Plotly charts.",
        6.0,
    )

    # ---------------------------------------------------------------- 3. concepts
    heading(doc, "3. Application of Course Concepts")
    para(
        doc,
        "Every concept is labelled in the source code with a `# Concept:` comment at the "
        "exact line where it is applied. The table maps each concept to where it is used.",
    )
    table(doc, ["Concept", "File / function", "How it is applied"], CONCEPTS, [1.3, 2.4, 2.8])

    # ---------------------------------------------------------------- 4. architecture
    heading(doc, "4. Architecture")
    para(
        doc,
        "The code is organised in layers. The three interfaces never touch files or pandas "
        "directly: they call `FinanceTracker` and the `analytics`, `alerts` and `visualize` "
        "modules, which rely on validated model objects and on `storage.py` for all disk "
        "access. Because the logic is independent of the interface, the same functions power "
        "the CLI, the GUI, the dashboard and the demo — and can be unit tested without a screen.",
    )
    figure(
        doc,
        ARCHITECTURE_PNG,
        "Figure 4 — Layered architecture: each layer only uses the layer below it.",
        6.2,
    )

    # ---------------------------------------------------------------- 5. challenges
    heading(doc, "5. Challenges and Solutions")
    para(doc, "These problems were met during development (see DEV_LOG.md for the full log).")
    table(doc, ["Challenge", "Solution"], CHALLENGES, [3.4, 3.1])

    # ---------------------------------------------------------------- 6. enhancements
    heading(doc, "6. Enhancements")
    bullets(
        doc,
        [
            "**Tkinter GUI** — four tabs (Transactions, Summary, Charts, Budgets & Goals), "
            "sortable Treeview, live search, embedded matplotlib charts, HiDPI support.",
            "**JSON and CSV storage** — JSON for the complete state, CSV for spreadsheets; both "
            "validated row by row, with missing-file and corrupted-file handling.",
            "**Budget alerts** — thresholds kept as constants (80 % / 100 %), sorted by severity, "
            "shown in every interface; goal milestones and missed deadlines are reported too.",
            "**Streamlit dashboard** — sidebar filters (source, dates, categories, kind), KPIs "
            "with month-over-month deltas, Plotly charts, heatmap, CSV upload and download.",
        ],
    )

    # ---------------------------------------------------------------- 7. testing
    heading(doc, "7. Testing")
    para(
        doc,
        f"The project has **{stats['tests']} automated pytest tests** with "
        f"**{stats['coverage']} % line coverage** of the `finance_tracker` package "
        "(the Tkinter and Streamlit event code is exercised by smoke tests instead). The tests "
        "cover normal use, invalid input (bad amounts, dates, kinds, empty names), missing and "
        "corrupted files, CSV files with missing columns or bad rows, every pandas calculation "
        "on a hand-checked dataset, and full CLI sessions driven by scripted answers. The GUI "
        "is built and clicked through on a real display, and the dashboard is rendered "
        "headlessly with Streamlit's `AppTest`. `ruff` enforces a consistent code style.",
    )

    # ---------------------------------------------------------------- 8. lessons
    heading(doc, "8. What I Learned")
    bullets(
        doc,
        [
            "Validating data once, at the boundary (`__post_init__`), keeps the rest of the "
            "code simple: everything after that point can trust its inputs.",
            "Choosing the right structure matters: a list keeps order, a dictionary gives fast "
            "lookup by name, and a set removes duplicates for free.",
            "pandas replaces dozens of lines of loops with one readable `groupby` or "
            "`pivot_table`, but empty data must be handled explicitly.",
            "Exception handlers are matched in order and by inheritance — a subtle source of bugs.",
            "Looking at the real output (screenshots, rendered charts) finds bugs that unit "
            "tests miss, such as the “$…$” math-text problem and the HiDPI layout issue.",
        ],
    )

    # ---------------------------------------------------------------- 9. future
    heading(doc, "9. Future Improvements")
    bullets(
        doc,
        [
            "Recurring transactions (rent, subscriptions) generated automatically each month.",
            "Import of real bank statements (OFX/QFX) with automatic category suggestions.",
            "A SQLite backend for larger histories and multi-user support.",
            "Spending forecasts (e.g. linear regression on monthly totals) in the dashboard.",
            "Packaging as an installable application with a `pyproject.toml` entry point.",
        ],
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main(argv: list[str] | None = None) -> int:
    """Command-line entry point.

    Args:
        argv: Command-line arguments (None = ``sys.argv[1:]``).

    Returns:
        Process exit code (0).
    """
    parser = argparse.ArgumentParser(description="Build the project report.")
    parser.add_argument("--skip-tests", action="store_true", help="do not run pytest")
    args = parser.parse_args(argv)
    print(f"✓ {build(args.skip_tests)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
