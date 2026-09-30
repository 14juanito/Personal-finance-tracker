"""Build the demo slide deck (deliverables/Demo_Slides.pptx) with python-pptx.

Eight 16:9 slides: title, overview, features, architecture, course concepts,
challenges, lessons learned, conclusion. Real screenshots from docs/screenshots/ and
the architecture diagram from docs/architecture.png are embedded; speaker notes hold
the talking points.

Usage:
    python tools/build_slides.py

Main elements:
    layout helpers (text, rich, card, badge, picture, title) and build.

Course concepts illustrated:
    Functions with keyword arguments, loops over lists of tuples.
"""

from __future__ import annotations

from pathlib import Path

from build_report import test_stats
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.shapes.autoshape import Shape
from pptx.shapes.picture import Picture
from pptx.slide import Slide
from pptx.util import Emu, Inches, Pt
from submission_config import COURSE, INSTITUTION, require_final_identity

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots"
ARCHITECTURE = ROOT / "docs" / "architecture.png"
OUTPUT = ROOT / "deliverables" / "Demo_Slides.pptx"

# "Money green" palette: deep green dominates, mint supports, gold is the accent.
DARK = RGBColor(0x1B, 0x43, 0x32)
GREEN = RGBColor(0x2D, 0x6A, 0x4F)
MINT = RGBColor(0x95, 0xD5, 0xB2)
TINT = RGBColor(0xEE, 0xF6, 0xF1)
GOLD = RGBColor(0xE9, 0xC4, 0x6A)
INK = RGBColor(0x21, 0x2B, 0x26)
MUTED = RGBColor(0x5C, 0x6B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB2, 0x3A, 0x48)
TITLE_FONT, BODY_FONT = "Cambria", "Calibri"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6


def background(slide: Slide, color: RGBColor) -> None:
    """Fill the slide background with a solid colour.

    Args:
        slide: Slide to modify.
        color: Background colour.
    """
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    content: str | list[str],
    size: int = 16,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = BODY_FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    italic: bool = False,
    spacing: int = 6,
) -> Shape:
    """Add a text box; a list becomes one paragraph per item.

    Args:
        slide: Slide to modify.
        x: Left position in inches.
        y: Top position in inches.
        w: Width in inches.
        h: Height in inches.
        content: One string or a list of paragraphs.
        size: Font size in points.
        color: Text colour.
        bold: Bold text.
        font: Font name.
        align: Horizontal alignment.
        anchor: Vertical alignment inside the box.
        italic: Italic text.
        spacing: Space after each paragraph, in points.

    Returns:
        The text box shape.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    lines = content if isinstance(content, list) else [content]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(spacing)
        run = paragraph.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.italic = Pt(size), bold, italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def rich(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    items: list[tuple[str, str]],
    size: int = 15,
    color: RGBColor = INK,
    spacing: int = 10,
) -> Shape:
    """Add paragraphs made of a bold lead-in followed by normal text.

    Args:
        slide: Slide to modify.
        x: Left position in inches.
        y: Top position in inches.
        w: Width in inches.
        h: Height in inches.
        items: ``(bold lead-in, rest of the sentence)`` pairs.
        size: Font size in points.
        color: Text colour.
        spacing: Space after each paragraph, in points.

    Returns:
        The text box shape.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, (lead, rest) in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(spacing)
        for part, bold in ((lead, True), (rest, False)):
            run = paragraph.add_run()
            run.text = part
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, BODY_FONT
            run.font.color.rgb = color
    return box


def card(slide: Slide, x: float, y: float, w: float, h: float, fill: RGBColor = TINT) -> Shape:
    """Add a rounded, tinted rectangle used as a content card.

    Args:
        slide: Slide to modify.
        x: Left position in inches.
        y: Top position in inches.
        w: Width in inches.
        h: Height in inches.
        fill: Card colour.

    Returns:
        The rectangle shape.
    """
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def badge(
    slide: Slide, x: float, y: float, glyph: str, fill: RGBColor = GREEN, size: float = 0.62
) -> Shape:
    """Add the deck's motif: a glyph in a filled circle.

    Args:
        slide: Slide to modify.
        x: Left position in inches.
        y: Top position in inches.
        glyph: One or two characters shown in the circle.
        fill: Circle colour.
        size: Diameter in inches.

    Returns:
        The circle shape.
    """
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    circle.line.fill.background()
    circle.shadow.inherit = False
    frame = circle.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = glyph
    run.font.size, run.font.bold, run.font.name = Pt(int(size * 26)), True, BODY_FONT
    run.font.color.rgb = WHITE if fill != GOLD else DARK
    return circle


def picture(
    slide: Slide, path: Path, x: float, y: float, max_w: float, max_h: float, border: bool = True
) -> Picture:
    """Insert an image scaled to fit a box, centred inside it.

    Args:
        slide: Slide to modify.
        path: Image file.
        x: Left of the box in inches.
        y: Top of the box in inches.
        max_w: Box width in inches.
        max_h: Box height in inches.
        border: Draw a thin border around the image.

    Returns:
        The picture shape.
    """
    with Image.open(path) as image:
        width, height = image.size
    scale = min(max_w / width, max_h / height)
    w, h = width * scale, height * scale
    left, top = x + (max_w - w) / 2, y + (max_h - h) / 2
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = RGBColor(0xC8, 0xD6, 0xCE)
        pic.line.width = Emu(12700)
    return pic


def title(slide: Slide, content: str, color: RGBColor = DARK) -> None:
    """Add the slide title in the heading font.

    Args:
        slide: Slide to modify.
        content: Title text.
        color: Title colour.
    """
    text(slide, MARGIN, 0.45, SLIDE_W - 2 * MARGIN, 0.9, content, 36, color, True, TITLE_FONT)


def build(output: Path = OUTPUT, skip_tests: bool = False) -> Path:
    """Create the eight slides and save the deck.

    Args:
        output: Destination ``.pptx`` file.
        skip_tests: Do not run pytest (the test figures then show "not run" markers).

    Returns:
        The path written.
    """
    student, submitted = require_final_identity()  # fail before the slow test run
    stats = test_stats(skip_tests)
    tests, coverage = stats["tests"], f"{stats['coverage']} %"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # 1 ── Title ---------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, DARK)
    badge(s, MARGIN, 1.2, "$", GOLD, 0.9)
    text(s, MARGIN, 2.3, 6.3, 1.6, "Personal Finance Tracker", 44, WHITE, True, TITLE_FONT)
    text(s, MARGIN, 3.95, 6.0, 1.0,
         "Track, analyse and visualise personal finances with Python, pandas, Tkinter and Streamlit",
         18, MINT)  # fmt: skip
    text(s, MARGIN, 5.4, 6.0, 1.0,
         [f"{COURSE} · {INSTITUTION}", f"{student} · {submitted}"],
         14, WHITE, spacing=4)  # fmt: skip
    picture(s, SHOTS / "dashboard_overview.png", 7.2, 1.0, 5.55, 5.5, border=False)
    s.notes_slide.notes_text_frame.text = (
        "Introduce yourself and the project: a Python personal finance tracker built for "
        "DATA 333. It has three interfaces — console, desktop and web dashboard — on top of "
        "one shared core."
    )

    # 2 ── Overview ------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, WHITE)
    title(s, "Project overview")
    rich(s, MARGIN, 1.6, 6.4, 4.8, [
        ("The problem. ", "Students rarely know where their money goes each month."),
        ("The solution. ", "Record income and expenses, categorise them, and turn them into "
                           "summaries, trends, budget alerts and savings-goal progress."),
        ("Three ways in. ", "A console menu, a Tkinter desktop app and a Streamlit dashboard — "
                            "plus a demo mode that runs everything non-interactively."),
        ("Realistic data. ", "A seeded generator creates 311 fictional transactions over six "
                             "months, so every run and screenshot is reproducible."),
    ], size=17, spacing=16)  # fmt: skip
    for index, (number, label) in enumerate(
        [
            ("311", "sample transactions"),
            ("3 + 1", "interfaces + demo mode"),
            (tests, "automated tests"),
            (coverage, "test coverage"),
        ]  # fmt: skip
    ):
        x, y = 7.6 + (index % 2) * 2.7, 1.7 + (index // 2) * 2.45
        card(s, x, y, 2.45, 2.15)
        text(s, x, y + 0.35, 2.45, 0.9, number, 40, GREEN, True, TITLE_FONT, PP_ALIGN.CENTER)
        text(s, x + 0.15, y + 1.35, 2.15, 0.6, label, 13, MUTED, align=PP_ALIGN.CENTER)
    s.notes_slide.notes_text_frame.text = (
        "Explain the goal and the scope: all course requirements plus the four enhancements "
        "(Tkinter GUI, JSON and CSV storage, budget alerts, Streamlit dashboard)."
    )

    # 3 ── Features ------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, WHITE)
    title(s, "Key features")
    features = [
        ("+", "Transactions", "Add, edit, delete, search and filter — with validated input."),
        ("%", "Analysis", "Monthly summary, category shares, rolling trend, top expenses."),
        ("!", "Budget alerts", "Warning at 80 %, exceeded at 100 % — shown instantly."),
        ("★", "Savings goals", "Targets, deadlines, contributions and 25/50/75/100 % milestones."),
        ("⇅", "JSON + CSV", "Atomic saves, CSV import/export, corrupted-file recovery."),
        ("▦", "Charts", "matplotlib PNGs in the GUI; interactive Plotly on the web."),
    ]  # fmt: skip
    for index, (glyph, head, body) in enumerate(features):
        col, row = index % 2, index // 2
        x, y = MARGIN + col * 3.55, 1.65 + row * 1.8
        card(s, x, y, 3.35, 1.6)
        badge(s, x + 0.2, y + 0.22, glyph)
        text(s, x + 1.0, y + 0.2, 2.2, 0.4, head, 16, DARK, True)
        text(s, x + 1.0, y + 0.6, 2.2, 0.95, body, 12, MUTED)
    picture(s, SHOTS / "gui_summary.png", 7.9, 1.65, 4.85, 3.3)
    text(s, 7.9, 5.1, 4.85, 0.4, "Tkinter GUI — Summary tab with colour-coded alerts", 11,
         MUTED, italic=True, align=PP_ALIGN.CENTER)  # fmt: skip
    s.notes_slide.notes_text_frame.text = (
        "Walk through the six feature groups. Point at the GUI screenshot: KPIs at the top, "
        "the monthly table, and alerts in red (exceeded) and orange (warning)."
    )

    # 4 ── Architecture --------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, WHITE)
    title(s, "Architecture: one core, three interfaces")
    picture(s, ARCHITECTURE, MARGIN, 1.5, 7.4, 5.4, border=False)
    rich(s, 8.4, 1.8, 4.3, 5.0, [
        ("Layers. ", "Interfaces call the application logic, which relies on validated "
                     "models and a single storage module."),
        ("Reuse. ", "The same analytics and chart functions feed the CLI, GUI, dashboard "
                    "and demo."),
        ("Testable. ", "Business rules live outside the UI, so they are unit tested without a screen."),
        ("Portable. ", "All paths come from pathlib in config.py — nothing hard-coded."),
    ], size=15, spacing=14)  # fmt: skip
    s.notes_slide.notes_text_frame.text = (
        "Explain the layers from top to bottom and why separating the UI from the logic made "
        "adding the GUI and dashboard easy."
    )

    # 5 ── Concepts ------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, WHITE)
    title(s, "Course concepts in the code")
    concepts = [
        ("Input / output", "cli.py — prompts, aligned tables"),
        ("Decisions", "alerts.budget_level — if / elif / else"),
        ("Loops", "re-prompt until a valid amount"),
        ("Functions", "parse_amount, analytics.*"),
        ("Files", "storage.py — json, csv, pathlib"),
        ("Exceptions", "custom classes, corrupted-file recovery"),
        ("Lists", "FinanceTracker.transactions"),
        ("Dictionaries", "budgets, goals, menu dispatch"),
        ("Sets", "unique categories, missing CSV columns"),
        ("OOP", "dataclasses + FinanceTracker class"),
        ("pandas", "pivot_table, groupby, rolling, pct_change"),
    ]
    for index, (concept, where) in enumerate(concepts):
        col, row = index % 3, index // 3
        x, y = MARGIN + col * 4.1, 1.6 + row * 1.3
        card(s, x, y, 3.9, 1.1)
        text(s, x + 0.25, y + 0.15, 3.5, 0.4, concept, 16, DARK, True)
        text(s, x + 0.25, y + 0.55, 3.5, 0.5, where, 12, MUTED)
    card(s, MARGIN + 2 * 4.1, 1.6 + 3 * 1.3, 3.9, 1.1, GREEN)
    text(s, MARGIN + 2 * 4.1 + 0.25, 1.6 + 3 * 1.3 + 0.15, 3.45, 0.85,
         "Each one is tagged with a  # Concept:  comment in the source.", 13, WHITE,
         anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    s.notes_slide.notes_text_frame.text = (
        "Every requirement from the brief is labelled in the code. Open one file, for example "
        "storage.py, to show a concept comment next to the try/except that recovers a "
        "corrupted JSON file."
    )

    # 6 ── Challenges ----------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, WHITE)
    title(s, "Challenges and solutions")
    challenges = [
        ("Money text turned into math",
         "\"$1,850 / $3,000\" rendered as LaTeX in matplotlib and Streamlit.",
         "Escape the dollar signs; regression tests."),
        ("Exception caught by the wrong handler",
         "Our StorageError is an OSError, so a generic handler re-wrapped it.",
         "Re-raise it first — handlers match in order."),
        ("Unreadable GUI on a HiDPI screen",
         "Tk scaled fonts but not row heights; chart text was tiny.",
         "Scale factor from the real font height; chart dpi."),
    ]  # fmt: skip
    text(s, 3.9, 1.55, 4.3, 0.4, "PROBLEM", 12, RED, True)
    text(s, 8.55, 1.55, 4.2, 0.4, "SOLUTION", 12, GREEN, True)
    for index, (head, problem, fix) in enumerate(challenges):
        y = 2.0 + index * 1.65
        badge(s, MARGIN, y + 0.35, str(index + 1), DARK, 0.55)
        text(s, MARGIN + 0.8, y + 0.2, 2.15, 1.2, head, 15, DARK, True, anchor=MSO_ANCHOR.MIDDLE)
        card(s, 3.75, y, 4.5, 1.4, RGBColor(0xFB, 0xEE, 0xF0))
        text(s, 3.95, y + 0.15, 4.1, 1.1, problem, 13, INK, anchor=MSO_ANCHOR.MIDDLE)
        card(s, 8.4, y, 4.33, 1.4)
        text(s, 8.6, y + 0.15, 3.95, 1.1, fix, 13, INK, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Pick one challenge to tell as a story — the dollar-sign bug is the most visual: it "
        "was only found by looking at the real rendered chart, not by the unit tests."
    )

    # 7 ── Lessons learned -----------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, WHITE)
    title(s, "What I learned")
    lessons = [
        ("✓", "Validate once, at the boundary",
         "__post_init__ checks every value, so the rest of the code can trust its data."),
        ("{}", "Pick the right data structure",
         "List for order, dictionary for lookup by name, set for uniqueness."),
        ("▤", "pandas replaces loops",
         "One groupby or pivot_table does what dozens of lines used to — but handle empty data."),
        ("◉", "Look at the real output",
         "Screenshots found bugs that the passing unit tests did not."),
    ]  # fmt: skip
    for index, (glyph, head, body) in enumerate(lessons):
        y = 1.6 + index * 1.35
        badge(s, MARGIN, y + 0.05, glyph, GREEN if index % 2 == 0 else GOLD)
        text(s, MARGIN + 0.9, y, 5.6, 0.4, head, 17, DARK, True)
        text(s, MARGIN + 0.9, y + 0.45, 5.6, 0.8, body, 13, MUTED)
    picture(s, SHOTS / "cli_summary.png", 7.5, 1.6, 5.25, 4.2)
    text(s, 7.5, 5.95, 5.25, 0.4, "CLI mode — monthly summary with category shares", 11,
         MUTED, italic=True, align=PP_ALIGN.CENTER)  # fmt: skip
    s.notes_slide.notes_text_frame.text = (
        "Share the main takeaways; connect each one to a concrete moment in the project."
    )

    # 8 ── Conclusion ----------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, DARK)
    text(s, MARGIN, 0.6, 12, 0.9, "Conclusion", 40, WHITE, True, TITLE_FONT)
    rich(s, MARGIN, 1.8, 6.2, 4.0, [
        ("Complete. ", "Every course requirement and all four enhancements are implemented."),
        ("Reliable. ", f"{tests} tests, {coverage} coverage, ruff-clean code, graceful error handling."),
        ("Reusable. ", "One tested core powers the console, desktop and web interfaces."),
        ("Next. ", "Recurring transactions, bank-statement import, spending forecasts."),
    ], size=17, color=WHITE, spacing=16)  # fmt: skip
    card(s, 7.3, 1.8, 5.43, 3.6, GREEN)
    text(s, 7.6, 2.0, 5.0, 0.4, "Try it", 18, GOLD, True)
    text(s, 7.6, 2.55, 5.0, 2.7, [
        "python main.py demo",
        "python main.py cli",
        "python main.py gui",
        "streamlit run finance_tracker/dashboard.py",
    ], 13, WHITE, font="Courier New", spacing=12)  # fmt: skip
    text(s, MARGIN, 6.3, 12, 0.6, "Thank you — questions?", 22, MINT, True, TITLE_FONT)
    s.notes_slide.notes_text_frame.text = "Summarise, show the four commands, and invite questions."

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


if __name__ == "__main__":
    print(f"✓ {build()}")
